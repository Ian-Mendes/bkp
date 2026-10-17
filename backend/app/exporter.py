import csv, io, re
from openpyxl import Workbook
from sqlalchemy import select, func, or_, case
from .models import Profile, TakeoutEntry

COLUMNS = ["username", "type", "profile_url", "followers_count", "following_count", "bio", "status"]
SORTS = {"followers_count": Profile.followers_count,
         "following_count": Profile.following_count,
         "username": Profile.username}
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FORMULA = ("=", "+", "-", "@", "\t", "\r")

def safe(v):
    """Remove caracteres ilegais e neutraliza injeção de fórmula em planilhas."""
    if isinstance(v, str):
        v = _ILLEGAL.sub("", v)
        if v[:1] in _FORMULA:
            v = "'" + v
    return v

def type(is_follower, is_following):
    """Relação do perfil com o dono do pedido, em uma coluna legível no backup."""
    if is_follower and is_following:
        return "ambos"
    if is_follower:
        return "seguidor"
    if is_following:
        return "seguindo"
    return ""

def row(profile, is_follower, is_following):
    return [safe(x) for x in (profile.username, type(is_follower, is_following),
                              profile.profile_url, profile.followers_count,
                              profile.following_count, profile.bio, profile.status)]

def query_profiles(order_id, sort_by="followers_count", direction="desc", q=None):
    """Linhas do pedido: `(Profile, is_follower, is_following)` derivados das entries.

    Agrega as relações por perfil via `sum(case(...))` (portável para SQLite;
    `bool_or` não existe no SQLite). Um perfil nas duas listas aparece uma única
    vez, com as duas flags em 1.
    """
    is_follower = func.sum(case((TakeoutEntry.relation == "follower", 1), else_=0))
    is_following = func.sum(case((TakeoutEntry.relation == "following", 1), else_=0))
    db_query = (
        select(Profile, is_follower.label("is_follower"), is_following.label("is_following"))
        .join(TakeoutEntry, TakeoutEntry.username == Profile.username)
        .where(TakeoutEntry.order_id == order_id)
        .group_by(Profile.id)
    )
    terms = [t.strip() for t in (q or "").split(",") if t.strip()]
    if terms:
        db_query = db_query.where(or_(*[Profile.bio.ilike(f"%{t}%") for t in terms]))
    col = SORTS.get(sort_by, Profile.followers_count)
    col = col.asc().nulls_last() if direction == "asc" else col.desc().nulls_last()
    return db_query.order_by(col, Profile.username)

def build_report(db, order):
    in_order = TakeoutEntry.order_id == order.id
    joined = TakeoutEntry.username == Profile.username

    total = db.scalar(
        select(func.count(func.distinct(Profile.id)))
        .join(TakeoutEntry, joined)
        .where(in_order)
    ) or 0
    by_status = dict(db.execute(
        select(Profile.status, func.count(func.distinct(Profile.id)))
        .join(TakeoutEntry, joined)
        .where(in_order)
        .group_by(Profile.status)
    ).all())
    coverage = {}
    for name, col in (("followers_count", Profile.followers_count),
                      ("following_count", Profile.following_count),
                      ("bio", Profile.bio),
                      ("profile_url", Profile.profile_url)):
        filled = db.scalar(
            select(func.count(func.distinct(Profile.id)))
            .join(TakeoutEntry, joined)
            .where(in_order, col.isnot(None))
        ) or 0
        coverage[name] = {"filled": filled, "total": total,
                          "pct": round(100 * filled / total, 1) if total else 0}

    follower = db.scalar(select(func.count(func.distinct(TakeoutEntry.username)))
                         .where(in_order, TakeoutEntry.relation == "follower")) or 0
    following = db.scalar(select(func.count(func.distinct(TakeoutEntry.username)))
                          .where(in_order, TakeoutEntry.relation == "following")) or 0
    both = db.scalar(select(func.count()).select_from(
        select(TakeoutEntry.username)
        .where(in_order)
        .group_by(TakeoutEntry.username)
        .having(func.count(func.distinct(TakeoutEntry.relation)) == 2)
        .subquery()
    )) or 0
    by_relationship = {"follower": follower - both, "following": following - both, "both": both}

    return {"order_id": str(order.id), "status": order.status, "total_profiles": total,
            "followers_in_file": order.followers_total, "following_in_file": order.following_total,
            "invalid": order.invalid_count, "duplicates": order.duplicate_count,
            "by_status": by_status, "by_relationship": by_relationship,
            "field_coverage": coverage, "export_info": order.info or {}}

def summary_rows(rep):
    rows = [("Resumo do pedido", ""), ("Perfis únicos", rep["total_profiles"]),
            ("Seguidores no arquivo", rep["followers_in_file"]),
            ("Seguindo no arquivo", rep["following_in_file"]),
            ("seguem você", rep["by_relationship"]["follower"]),
            ("você segue", rep["by_relationship"]["following"]),
            ("ambos", rep["by_relationship"]["both"]),
            ("Registros inválidos", rep["invalid"]), ("Duplicados removidos", rep["duplicates"])]
    for s, n in rep["by_status"].items():
        rows.append((f"status: {s}", n))
    for f, c in rep["field_coverage"].items():
        rows.append((f"preenchido: {f}", f"{c['filled']}/{c['total']} ({c['pct']}%)"))
    info = rep["export_info"]
    for k, c in (info.get("coverage") or {}).items():
        rows.append((f"cobertura {k}", f"{c['found']}/{c['expected']} ({round(c['ratio'] * 100, 1)}%)"))
    for w in info.get("warnings", []):
        rows.append(("aviso", w))
    return rows

def to_csv(rows):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COLUMNS)
    for profile, is_follower, is_following in rows:
        w.writerow(["" if v is None else v for v in row(profile, is_follower, is_following)])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")

def to_xlsx(rows, rep):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Perfis")
    ws.append(COLUMNS)
    for profile, is_follower, is_following in rows:
        ws.append(row(profile, is_follower, is_following))
    rs = wb.create_sheet("Resumo")
    for r in summary_rows(rep):
        rs.append(list(r))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
