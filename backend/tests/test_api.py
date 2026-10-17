import csv, io, os
import openpyxl
from sqlalchemy import select, func
from app import main
from app.models import Profile, TakeoutEntry
from app.services import run_enrichment
from tests.helpers import make_zip, register, upload, OLD, NEW

def _setup(env):
    client, Session, tmp = env
    h = register(client)
    p = tmp / "e.zip"
    make_zip(p, followers=[("ana_silva", OLD), ("bob", NEW)], following=[("ana_silva", OLD), ("carla", NEW)])
    r = upload(client, p, h, expected_followers="10")
    assert r.status_code == 200, r.text
    return client, Session, h, r.json()

def test_upload_e_cobertura(env):
    client, Session, h, body = _setup(env)
    assert body["followers"] == 2 and body["following"] == 2
    assert "followers_incomplete" in body["warnings"]

def test_zip_apagado_apos_processar(env):
    _setup(env)
    assert os.listdir(main.UPLOAD_DIR) == []

def test_arquivo_invalido(env):
    client, _, _ = env
    h = register(client)
    r = client.post("/orders", headers=h, files={"file": ("x.zip", b"nao e zip", "application/zip")})
    assert r.status_code == 400

def test_perfil_so_guarda_campos_permitidos():
    cols = {c.name for c in Profile.__table__.columns}
    assert cols == {"id", "username", "profile_url", "followers_count", "following_count", "bio", "status", "enrich_source", "enriched_at"}
    entries = {c.name for c in TakeoutEntry.__table__.columns}
    assert entries == {"id", "order_id", "username", "relation"}

def test_relacao_ambos_gera_duas_entries(env):
    """Usuário nas duas listas gera 2 entries (follower + following) num só perfil."""
    client, Session, tmp = env
    h = register(client)
    p = tmp / "e.zip"
    make_zip(p, followers=[("ana_silva", OLD), ("bob", NEW)], following=[("ana_silva", OLD), ("carla", NEW)])
    upload(client, p, h)
    with Session() as db:
        rel = {(e.username, e.relation) for e in db.scalars(select(TakeoutEntry))}
        assert rel == {("ana_silva", "follower"), ("ana_silva", "following"),
                       ("bob", "follower"), ("carla", "following")}
        assert db.scalar(select(func.count()).select_from(Profile)) == 3

def test_worker_processa_tudo(env):
    client, Session, h, body = _setup(env)
    run_enrichment(body["order_id"])
    r = client.get(f"/orders/{body['order_id']}", headers=h).json()
    assert r["status"] == "done"
    assert r["profiles_by_status"] == {"not_enriched": 3}

def test_csv_ordenacao_null_zero_e_formula(env):
    client, Session, h, body = _setup(env)
    oid = body["order_id"]
    run_enrichment(oid)
    with Session() as db:
        ps = {x.username: x for x in db.scalars(select(Profile))}
        ps["ana_silva"].followers_count = 500
        ps["ana_silva"].bio = '=HYPERLINK("http://x")'
        ps["bob"].followers_count = 0
        ps["bob"].bio = "Designer em Natal"
        db.commit()
    r = client.get(f"/orders/{oid}/export.csv?sort_by=followers_count&direction=desc", headers=h)
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert list(rows[0].keys()) == ["username", "tipo", "profile_url", "followers_count", "following_count", "bio", "status"]
    assert [x["username"] for x in rows] == ["ana_silva", "bob", "carla"]
    assert [x["tipo"] for x in rows] == ["ambos", "seguidor", "seguindo"]
    assert rows[0]["bio"].startswith("'=")          # fórmula neutralizada
    assert rows[1]["followers_count"] == "0"        # zero de verdade
    assert rows[2]["followers_count"] == ""         # não consultado = vazio

    r = client.get(f"/orders/{oid}/export.csv?q=designer", headers=h)
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert [x["username"] for x in rows] == ["bob"]

def test_relatorio_e_xlsx(env):
    client, Session, h, body = _setup(env)
    oid = body["order_id"]
    run_enrichment(oid)
    rep = client.get(f"/orders/{oid}/report", headers=h).json()
    assert rep["total_profiles"] == 3
    assert rep["by_relationship"] == {"follower": 1, "following": 1, "both": 1}
    assert set(rep["field_coverage"]) == {"followers_count", "following_count", "bio", "profile_url"}
    assert rep["field_coverage"]["followers_count"]["filled"] == 0
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/orders/{oid}/export.xlsx", headers=h).content))
    assert wb.sheetnames == ["Perfis", "Resumo"]
    assert wb["Perfis"].max_row == 4


def test_limite_upload_autenticado(env, monkeypatch):
    client, _, _ = env
    assert client.get("/orders/upload-config").status_code == 401
    h = register(client)
    monkeypatch.setattr(main, "MAX_UPLOAD_MB", 7)
    response = client.get("/orders/upload-config", headers=h)
    assert response.status_code == 200
    assert response.json() == {"max_upload_bytes": 7 * 1024 * 1024}
    assert response.headers["Cache-Control"] == "no-store"


def test_profiles_paginado(env):
    client, Session, h, body = _setup(env)
    oid = body["order_id"]
    r = client.get(f"/orders/{oid}/profiles?limit=2&offset=0", headers=h)
    assert r.status_code == 200
    data = r.json()
    assert data["total"] == 3
    assert data["limit"] == 2 and data["offset"] == 0
    assert [p["username"] for p in data["profiles"]] == ["ana_silva", "bob"]
    r2 = client.get(f"/orders/{oid}/profiles?limit=2&offset=2", headers=h).json()
    assert r2["total"] == 3
    assert [p["username"] for p in r2["profiles"]] == ["carla"]


def test_profiles_filtro_tipo(env):
    client, Session, h, body = _setup(env)
    oid = body["order_id"]
    assert [p["username"] for p in client.get(f"/orders/{oid}/profiles?tipo=ambos", headers=h).json()["profiles"]] == ["ana_silva"]
    assert [p["username"] for p in client.get(f"/orders/{oid}/profiles?tipo=seguidor", headers=h).json()["profiles"]] == ["bob"]
    assert [p["username"] for p in client.get(f"/orders/{oid}/profiles?tipo=seguindo", headers=h).json()["profiles"]] == ["carla"]


def test_profiles_q(env):
    client, Session, h, body = _setup(env)
    oid = body["order_id"]
    with Session() as db:
        ps = {x.username: x for x in db.scalars(select(Profile))}
        ps["bob"].bio = "Designer em Natal"
        db.commit()
    r = client.get(f"/orders/{oid}/profiles?q=designer", headers=h).json()
    assert r["total"] == 1
    assert r["profiles"][0]["username"] == "bob"


def test_profiles_isolamento(env):
    client, Session, h, body = _setup(env)
    oid = body["order_id"]
    hb = register(client, "bia@empresa.com.br")
    assert client.get(f"/orders/{oid}/profiles", headers=hb).status_code == 404


def test_profiles_null_nao_zero(env):
    client, Session, h, body = _setup(env)
    oid = body["order_id"]
    run_enrichment(oid)   # NullProvider → not_enriched, contadores ficam NULL
    r = client.get(f"/orders/{oid}/profiles", headers=h).json()
    assert r["total"] == 3
    for p in r["profiles"]:
        assert p["status"] == "not_enriched"
        assert p["followers_count"] is None
        assert p["following_count"] is None
        assert p["bio"] is None
